"""
DEIMOS — Offer Collector v2 (Playwright)
Scraping da Meta Ad Library sem token.
"""

import asyncio
import re
import sys
from datetime import date, datetime
from typing import Optional

from playwright.async_api import async_playwright, Page, Browser, BrowserContext
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Garante saída UTF-8 no Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# ── Constantes ────────────────────────────────────────────────────────────────

AD_LIBRARY_BASE = "https://www.facebook.com/ads/library/"

# ── Estilos Excel ─────────────────────────────────────────────────────────────

FILL_HEADER = PatternFill("solid", fgColor="1F2D3D")
FILL_ALT    = PatternFill("solid", fgColor="F2F4F6")
FILL_WHITE  = PatternFill("solid", fgColor="FFFFFF")
FILL_RED    = PatternFill("solid", fgColor="FFCCCC")
FILL_ORANGE = PatternFill("solid", fgColor="FFE0B2")
FILL_YELLOW = PatternFill("solid", fgColor="FFF9C4")
FILL_GRAY   = PatternFill("solid", fgColor="EEEEEE")
FILL_GREEN  = PatternFill("solid", fgColor="D4EDDA")

FONT_HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
FONT_TITLE  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
FONT_BODY   = Font(name="Calibri", size=10)
FONT_LINK   = Font(name="Calibri", size=10, color="1155CC", underline="single")

COLUMNS = [
    ("collected_at",     "Coletado em",       18),
    ("keyword",          "Keyword",           22),
    ("page_name",        "Anunciante",        30),
    ("active_ads_count", "Ads Ativos",        12),
    ("has_duplicates",   "Duplicatas",        12),
    ("first_ad_date",    "Data início",       16),
    ("sample_copy",      "Anúncio",           52),
    ("offer",            "Oferta detectada",  32),
    ("page_ad_url",      "Link Ad Library",   14),
    ("signal_strength",  "Sinal",             12),
]

# ── Detecção de oferta ────────────────────────────────────────────────────────

_OFFER_PATTERNS = [
    r'R\$\s*[\d.,]+(?:\s*(?:mil|k))?',
    r'\d+\s*[xX]\s*(?:de\s*)?R?\$?\s*[\d.,]+',
    r'\d+\s*(?:parcelas?)',
    r'\d+\s*%\s*(?:off|de\s*desconto|desconto)',
    r'(?:grátis|gratuito|free|de\s*graça)',
    r'(?:acesso\s+(?:vitalício|perpétuo|imediato))',
    r'(?:garantia\s+(?:de\s*)?\d+\s*dias?)',
    r'(?:bônus|bonus)\s*[:\-]?\s*[^\n.!]{5,40}',
]


def extract_offer(text: str) -> str:
    if not text:
        return ""
    found = []
    for pattern in _OFFER_PATTERNS:
        for m in re.findall(pattern, text, re.IGNORECASE):
            m = m.strip()
            if m and m not in found:
                found.append(m)
        if len(found) >= 3:
            break
    return " | ".join(found[:3])


def classify_signal(count: int) -> str:
    if count >= 20:
        return "FORTE"
    if count >= 10:
        return "MÉDIO"
    if count >= 5:
        return "FRACO"
    return "RUÍDO"


def _signal_fill(value: str) -> PatternFill:
    if "FORTE" in value:
        return FILL_RED
    if "MÉDIO" in value:
        return FILL_ORANGE
    if "FRACO" in value:
        return FILL_YELLOW
    return FILL_GRAY


# ── Scraping ──────────────────────────────────────────────────────────────────

async def _dismiss_popups(page: Page) -> None:
    """Fecha banners de cookies e popups comuns."""
    selectors = [
        'button:has-text("Allow all cookies")',
        'button:has-text("Aceitar todos")',
        'button:has-text("Allow essential and optional cookies")',
        '[aria-label="Close"]',
        '[data-testid="cookie-policy-manage-dialog-accept-button"]',
    ]
    for sel in selectors:
        try:
            btn = page.locator(sel).first
            if await btn.is_visible(timeout=800):
                await btn.click()
                await page.wait_for_timeout(600)
        except Exception:
            pass


async def _scroll_to_load(page: Page, max_scrolls: int) -> None:
    """Rola a página para forçar carregamento de mais resultados."""
    for _ in range(max_scrolls):
        prev_height = await page.evaluate("document.body.scrollHeight")
        await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        await page.wait_for_timeout(2200)
        new_height = await page.evaluate("document.body.scrollHeight")
        if new_height == prev_height:
            break


def _parse_ad_count(text: str) -> int:
    """Extrai número de ads ativos do texto do card."""
    # "32 anúncios ativos" | "1 anúncio ativo" | "32 active ads"
    m = re.search(r'([\d.,]+)\s+(?:anúncios?\s+ativos?|active\s+ads?)', text, re.IGNORECASE)
    if m:
        return int(m.group(1).replace(".", "").replace(",", ""))
    return 0


def _parse_date(text: str) -> str:
    """Extrai a primeira data encontrada no texto."""
    patterns = [
        r'\d{1,2}\s+de\s+\w+\s+de\s+\d{4}',   # 3 de janeiro de 2024
        r'\d{1,2}/\d{2}/\d{4}',                  # 03/01/2024
        r'\w+\s+\d{1,2},\s+\d{4}',               # January 3, 2024
        r'\d{4}-\d{2}-\d{2}',                     # 2024-01-03
    ]
    for p in patterns:
        m = re.search(p, text, re.IGNORECASE)
        if m:
            return m.group(0)
    return ""


async def _extract_cards(page: Page, keyword: str) -> list[dict]:
    """
    Extrai dados dos cards de resultado da Ad Library.
    Agrupa por anunciante (page_id) e soma contagem de ads.

    Estrutura real do card:
      - Link facebook.com/PAGE_ID/ → nome do anunciante
      - "Veiculação iniciada em X de X de XXXX"
      - "N anúncios usam esse criativo e esse texto"
      - Texto do anúncio (quando disponível)
    """
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")

    # page_id → dados agregados
    advertisers: dict[str, dict] = {}

    # Links para páginas de anunciantes: facebook.com/NUMERIC_ID/
    links = await page.query_selector_all('a[href*="facebook.com/"]')

    for link in links:
        try:
            href = await link.get_attribute("href") or ""
            # Filtrar apenas links com ID numérico
            m = re.search(r'facebook\.com/(\d{5,})', href)
            if not m:
                continue

            page_id = m.group(1)
            page_name = (await link.inner_text() or "").strip()
            if not page_name or len(page_name) < 2:
                continue

            # Pegar texto do card subindo na DOM
            card_text: str = await link.evaluate("""el => {
                let node = el;
                for (let i = 0; i < 12; i++) {
                    node = node.parentElement;
                    if (!node) break;
                    const t = node.innerText || '';
                    if (t.includes('Veiculação iniciada') || t.includes('anúncios usam')) {
                        return t.substring(0, 800);
                    }
                }
                return '';
            }""")

            if not card_text:
                continue

            # Quantos ads usam esse criativo
            dupe_m = re.search(
                r'(\d+)\s+anúncios?\s+usam\s+esse\s+criativo',
                card_text, re.IGNORECASE
            )
            creative_count = int(dupe_m.group(1)) if dupe_m else 1

            # Library ID do anúncio → URL direta do ad
            lib_m = re.search(r'Identificação da biblioteca[:\s]+(\d+)', card_text)
            library_id = lib_m.group(1) if lib_m else ""

            # Data de início
            first_date = _parse_date(card_text)

            # Copy: linhas longas que não são metadados
            _META = re.compile(
                r'Veiculação iniciada|anúncios usam|Identificação da biblioteca'
                r'|Abrir menu|Ver resumo|Patrocinado|Plataformas|Ativo|Inativo'
                r'|^\d+:\d+\s*/\s*\d+:\d+$',
                re.IGNORECASE
            )
            copy_lines = [
                l.strip() for l in card_text.splitlines()
                if len(l.strip()) > 25
                and not _META.search(l.strip())
                and l.strip() != page_name
            ]
            sample_copy = " ".join(copy_lines[:4])[:350]

            # Agregar por anunciante
            if page_id not in advertisers:
                advertisers[page_id] = {
                    "page_name":    page_name,
                    "page_id":      page_id,
                    "total_ads":    0,
                    "max_creative": 0,
                    "best_lib_id":  library_id,
                    "first_date":   first_date,
                    "sample_copy":  sample_copy,
                    "copies":       [],
                }

            adv = advertisers[page_id]
            adv["total_ads"] += creative_count
            # Guardar library_id do criativo mais duplicado
            if creative_count > adv["max_creative"]:
                adv["max_creative"] = creative_count
                adv["best_lib_id"]  = library_id
            if sample_copy and sample_copy not in adv["copies"]:
                adv["copies"].append(sample_copy)
            if first_date and (not adv["first_date"] or first_date < adv["first_date"]):
                adv["first_date"] = first_date

        except Exception:
            continue

    # Montar resultados
    results: list[dict] = []
    for page_id, adv in advertisers.items():
        lib_id = adv.get("best_lib_id", "")
        ad_url = (
            f"{AD_LIBRARY_BASE}?id={lib_id}"
            if lib_id else
            f"{AD_LIBRARY_BASE}?active_status=active&ad_type=all&country=BR&view_all_page_id={page_id}"
        )

        best_copy = adv["copies"][0] if adv["copies"] else ""
        total_ads = adv["total_ads"]
        # Duplicata: mesmo criativo rodando em 2+ campanhas OU anunciante com 3+ ads total
        has_duplicates = adv["max_creative"] >= 2 or total_ads >= 3

        results.append({
            "collected_at":     now_str,
            "keyword":          keyword,
            "page_name":        adv["page_name"],
            "active_ads_count": total_ads,
            "has_duplicates":   "Sim" if has_duplicates else "Não",
            "first_ad_date":    adv["first_date"],
            "sample_copy":      best_copy,
            "offer":            extract_offer(best_copy),
            "page_ad_url":      ad_url,
            "signal_strength":  classify_signal(total_ads),
        })

    results.sort(key=lambda x: x["active_ads_count"], reverse=True)
    return results


async def scrape_keyword(
    page: Page,
    keyword: str,
    country: str = "BR",
    max_scrolls: int = 5,
) -> list[dict]:
    """Pesquisa uma keyword na Ad Library e retorna lista de anunciantes."""
    url = (
        f"{AD_LIBRARY_BASE}"
        f"?active_status=active&ad_type=all&country={country}"
        f"&q={keyword}&search_type=keyword_unordered&media_type=all"
    )

    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=35000)
    except Exception:
        await page.goto(url, timeout=35000)

    await page.wait_for_timeout(3000)
    await _dismiss_popups(page)
    await _scroll_to_load(page, max_scrolls)

    return await _extract_cards(page, keyword)


CHROME_PATH = "C:/Program Files/Google/Chrome/Application/chrome.exe"


async def create_browser_context(playwright) -> tuple:
    """Cria browser e context com configurações anti-detecção básicas."""
    browser: Browser = await playwright.chromium.launch(
        executable_path=CHROME_PATH,
        headless=False,
        args=[
            "--disable-blink-features=AutomationControlled",
            "--no-sandbox",
        ],
    )
    context: BrowserContext = await browser.new_context(
        locale="pt-BR",
        timezone_id="America/Sao_Paulo",
        viewport={"width": 1366, "height": 768},
        user_agent=(
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/124.0.0.0 Safari/537.36"
        ),
    )
    # Remove webdriver flag
    await context.add_init_script(
        "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
    )
    return browser, context


# ── Excel ─────────────────────────────────────────────────────────────────────

def _thin_border() -> Border:
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def export_to_excel(
    observations: list[dict],
    output_path: str,
    total_keywords: int,
    keyword_stats: Optional[list[dict]] = None,
) -> None:
    wb = openpyxl.Workbook()
    _build_obs_sheet(wb.active, observations, total_keywords)
    if keyword_stats:
        _build_summary_sheet(wb.create_sheet("Resumo"), observations, keyword_stats)
    wb.save(output_path)


def _build_obs_sheet(ws, observations: list[dict], total_keywords: int) -> None:
    ws.title = "Offer Observations"
    total = len(observations)
    today = date.today().strftime("%d/%m/%Y")

    # Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    t = ws.cell(row=1, column=1)
    t.value = f"DEIMOS · Offer Observations · {today} · {total} players · {total_keywords} keywords"
    t.font = FONT_TITLE
    t.fill = FILL_HEADER
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Cabeçalhos
    for col_idx, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=label)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 20

    col_keys = [c[0] for c in COLUMNS]

    for row_idx, obs in enumerate(observations, start=3):
        base_fill = FILL_WHITE if row_idx % 2 == 1 else FILL_ALT

        for col_idx, key in enumerate(col_keys, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = FONT_BODY
            cell.border = _thin_border()
            cell.alignment = Alignment(vertical="top", wrap_text=(key == "sample_copy"))

            if key == "page_ad_url":
                url = obs.get(key, "")
                cell.value = url
                cell.fill = base_fill
                continue

            if key == "signal_strength":
                val = obs.get(key, "")
                cell.value = val
                cell.fill = _signal_fill(val)
                cell.alignment = Alignment(horizontal="center", vertical="top")
                continue

            if key == "has_duplicates":
                val = obs.get(key, "")
                cell.value = val
                cell.fill = FILL_GREEN if val == "Sim" else base_fill
                cell.alignment = Alignment(horizontal="center", vertical="top")
                continue

            if key == "active_ads_count":
                cell.value = obs.get(key, 0)
                cell.fill = base_fill
                cell.alignment = Alignment(horizontal="center", vertical="top")
                continue

            cell.value = obs.get(key, "")
            cell.fill = base_fill

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}2"


def _build_summary_sheet(ws, observations: list[dict], keyword_stats: list[dict]) -> None:
    ws.title = "Resumo"

    signal_order = ["FORTE", "MÉDIO", "FRACO", "RUÍDO"]
    counts = {s: 0 for s in signal_order}
    for obs in observations:
        sig = obs.get("signal_strength", "RUÍDO")
        counts[sig] = counts.get(sig, 0) + 1

    # Título bloco 1
    ws.merge_cells("A1:C1")
    t = ws.cell(row=1, column=1, value="Resumo por Força de Sinal")
    t.font = FONT_TITLE
    t.fill = FILL_HEADER
    t.alignment = Alignment(horizontal="center")

    for col, label in enumerate(["Sinal", "Players", "% do Total"], start=1):
        c = ws.cell(row=2, column=col, value=label)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")

    total = len(observations) or 1
    signal_fills = {"FORTE": FILL_RED, "MÉDIO": FILL_ORANGE, "FRACO": FILL_YELLOW, "RUÍDO": FILL_GRAY}
    for row_idx, sig in enumerate(signal_order, start=3):
        cnt = counts[sig]
        ws.cell(row=row_idx, column=1, value=sig).fill = signal_fills[sig]
        ws.cell(row=row_idx, column=2, value=cnt)
        pct = ws.cell(row=row_idx, column=3, value=round(cnt / total * 100, 1))

    # Título bloco 2
    start_row = 10
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=4)
    t2 = ws.cell(row=start_row, column=1, value="Stats por Keyword")
    t2.font = FONT_TITLE
    t2.fill = FILL_HEADER
    t2.alignment = Alignment(horizontal="center")

    headers = ["Keyword", "Players encontrados", "Novos únicos", "Duplicatas"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row + 1, column=col_idx, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")

    for r, stat in enumerate(keyword_stats, start=start_row + 2):
        fill = FILL_WHITE if r % 2 == 0 else FILL_ALT
        ws.cell(row=r, column=1, value=stat.get("keyword", "")).fill = fill
        ws.cell(row=r, column=2, value=stat.get("players", 0)).fill = fill
        ws.cell(row=r, column=3, value=stat.get("novos_unicos", 0)).fill = fill
        ws.cell(row=r, column=4, value=stat.get("duplicatas", 0)).fill = fill

    for col, width in [(1, 28), (2, 20), (3, 14), (4, 12)]:
        ws.column_dimensions[get_column_letter(col)].width = width
